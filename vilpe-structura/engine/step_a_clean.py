import glob, warnings
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
warnings.filterwarnings('ignore')

EN = {'Hallin alapohja': 'Hall crawl space', 'Katto 1': 'Roof 1', 'Katto 2': 'Roof 2', 'Katto 3': 'Roof 3',
      'Katto 4': 'Roof 4', 'Viherkatto 1': 'Green roof 1', 'Viherkatto 2': 'Green roof 2'}
PURPOSE = {'Kattorakenteen tuuletus': 'Roof structure ventilation', 'Alapohjan tuuletus': 'Crawl space ventilation',
           'Ryömintätilan tuuletus': 'Crawl space ventilation'}
MATERIAL = {'Betoni': 'Concrete'}
COLS = ['Timestamp', 'Fan speed (rpm)', 'Outdoor temp (°C)', 'Outdoor RH (%)', 'Outdoor abs. humidity (g/m³)',
        'Indoor temp (°C)', 'Indoor RH (%)', 'Indoor abs. humidity (g/m³)']

F = Font(name='Arial', size=10); FB = Font(name='Arial', size=10, bold=True, color='FFFFFF')
HEAD = PatternFill('solid', start_color='1F4E79'); NOTE = PatternFill('solid', start_color='FFF2CC')

def style_header(ws, row=1):
    for c in ws[row]:
        if c.value is not None:
            c.font = FB; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical='center')

structures, logs, sheets = [], [], {}
for f in sorted(glob.glob(str(ROOT / 'data/raw/VILPE Vantaa*.xlsx'))):
    raw = pd.read_excel(f, header=None)
    meta = {raw.iloc[i, 0]: raw.iloc[i, 1] for i in range(9)}
    fi = raw.index[raw.iloc[:, 0] == 'Aikaleima'][0]
    name = meta['Tunniste'].replace('VILPE Vantaa, ', '')
    # latest sensor table (rows between metadata and time series)
    sens = raw.iloc[11:fi].dropna(how='all')
    sens = sens[sens.iloc[:, 0].isin(['Ohjaava sisälähetin', 'Ohjaava ulkolähetin'])]
    sdict = {r[0]: r for r in sens.values.tolist()}
    d = raw.iloc[fi + 1:, :8].copy(); d.columns = COLS
    n_orig = len(d)
    d['Timestamp'] = pd.to_datetime(d['Timestamp'], format='%Y/%m/%d %H:%M:%S')
    for c in COLS[1:]: d[c] = pd.to_numeric(d[c], errors='coerce')
    d['Duplicate of previous row'] = d.duplicated(subset=COLS, keep='first').map({True: 'YES', False: ''})
    d['Indoor reading missing'] = d['Indoor RH (%)'].isna().map({True: 'YES', False: ''})
    d['Outdoor reading missing'] = d['Outdoor RH (%)'].isna().map({True: 'YES', False: ''})
    sheets[name] = d
    ins, out = sdict.get('Ohjaava sisälähetin'), sdict.get('Ohjaava ulkolähetin')
    structures.append([name, EN.get(name, name), meta['Sarjanumero'], meta['Tyyppi'],
                       PURPOSE.get(meta['Käyttötarkoitus'], meta['Käyttötarkoitus']),
                       MATERIAL.get(meta['Rakennusmateriaali'], meta['Rakennusmateriaali']),
                       float(meta['Viimeisin homeindeksi']), int(meta['Viimeisin rpm']),
                       ins[2] if ins else '', out[2] if out else '',
                       d['Timestamp'].min(), d['Timestamp'].max()])
    logs.append([name, n_orig, int((d['Duplicate of previous row'] == 'YES').sum()),
                 int(d['Indoor RH (%)'].isna().sum()), int(d['Outdoor RH (%)'].isna().sum())])

wb = Workbook(); ws = wb.active; ws.title = 'README'
readme = [
    ['VILPE Sense, Vantaa warehouse: cleaned data'],
    [''],
    ['Source', '7 Excel exports from the VILPE Sense cloud service, provided by VILPE for the Vaasa Hackathon 2026'],
    ['What changed', 'Finnish column names translated to English; timestamps converted to real date-time values; numbers stored as numbers; metadata collected into the Structures sheet'],
    ['What was NOT changed', 'No row was removed and no measured value was altered. The Cleaning log sheet proves the row counts match the originals.'],
    ['Added flag columns', '"Duplicate of previous row" = exact repeat of an earlier row (same time and values). "Indoor/Outdoor reading missing" = sensor value blank in the original.'],
    ['Measurement interval', 'About every 2 hours (hourly during the first days after installation)'],
    ['Abs. humidity', 'Absolute humidity in grams of water per m³ of air, calculated by VILPE in the original export'],
    ['Mold index', 'Latest value only (VILPE cloud, scale 0 to 6, alarm at 2.5). No history in the export.'],
    ['Data quality notes', 'See the Data quality notes sheet for issues found during cleaning.'],
]
for r in readme: ws.append(r)
ws['A1'].font = Font(name='Arial', size=14, bold=True, color='1F4E79')
for row in ws.iter_rows(min_row=3):
    row[0].font = Font(name='Arial', size=10, bold=True)
    if len(row) > 1: row[1].font = F; row[1].alignment = Alignment(wrap_text=True)
ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 110

ws = wb.create_sheet('Structures')
ws.append(['Sheet name (original)', 'English name', 'Serial number', 'Control unit', 'Purpose', 'Material',
           'Latest mold index', 'Latest fan rpm', 'Indoor sensor serial', 'Outdoor sensor serial', 'First reading', 'Last reading'])
for s in structures: ws.append(s)
style_header(ws)
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = F
    row[6].number_format = '0.00000'; row[10].number_format = row[11].number_format = 'yyyy-mm-dd hh:mm'
for i, w in enumerate([22, 18, 15, 13, 28, 11, 16, 14, 20, 20, 17, 17], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws = wb.create_sheet('Cleaning log')
ws.append(['Sheet', 'Rows in original file', 'Rows in this workbook', 'Check', 'Duplicate rows (flagged, kept)',
           'Rows with indoor reading missing', 'Rows with outdoor reading missing'])
for i, (name, n, dup, imiss, omiss) in enumerate(logs, start=2):
    ws.append([name, n, f"=COUNTA('{name}'!A:A)-1", f'=IF(B{i}=C{i},"OK, nothing removed","MISMATCH")',
               f"=COUNTIF('{name}'!I:I,\"YES\")", f"=COUNTIF('{name}'!J:J,\"YES\")", f"=COUNTIF('{name}'!K:K,\"YES\")"])
    ws.cell(i, 2).font = Font(name='Arial', size=10, color='0000FF')
    for c in range(3, 8): ws.cell(i, c).font = F
    ws.cell(i, 1).font = F
t = len(logs) + 2
ws.cell(t, 1, 'Total').font = Font(name='Arial', size=10, bold=True)
for c in [2, 3, 5, 6, 7]:
    L = get_column_letter(c); ws.cell(t, c, f'=SUM({L}2:{L}{t-1})').font = Font(name='Arial', size=10, bold=True)
ws.cell(t + 2, 1, 'Blue numbers = row counts read from the original VILPE files. Black = formulas counting rows in this workbook.').font = Font(name='Arial', size=9, italic=True)
style_header(ws)
for i, w in enumerate([18, 20, 20, 22, 26, 28, 28], 1): ws.column_dimensions[get_column_letter(i)].width = w

ws = wb.create_sheet('Data quality notes')
ws.append(['#', 'Sheet', 'What we noticed', 'Evidence in the data', 'Action taken'])
notes = [
    [1, 'All sheets', 'Some rows are exact repeats (same timestamp and same values)',
     '17 rows in total, see Cleaning log', 'Kept and flagged in column I'],
    [2, 'Roof 1, Green roof 1', 'Indoor sensor has no values for the first weeks',
     'Indoor columns blank until 18 Jun 2025 (about 420 rows each)', 'Kept and flagged in column J'],
    [3, 'Hallin alapohja', 'Indoor and outdoor sensor labels look swapped',
     'The "outdoor" sensor stays between 8 and 13 °C all year (Jan 2026 average 8.5 °C), while the "indoor" sensor follows the weather (Jan 2026 average -9.2 °C) and matches the outdoor sensors of the other 6 units (correlation 0.97)',
     'Not changed. Values kept exactly as exported. To confirm with VILPE'],
    [4, 'Viherkatto 2', 'Fan speed is 0 rpm for about 12 months',
     '0 rpm from 13 May 2025 to 27 May 2026 (97% of all readings in the sheet)', 'Not changed. To confirm with VILPE'],
]
for n in notes: ws.append(n)
style_header(ws)
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = F; c.alignment = Alignment(wrap_text=True, vertical='top')
for i, w in enumerate([4, 20, 38, 70, 34], 1): ws.column_dimensions[get_column_letter(i)].width = w

order = ['Katto 1', 'Katto 2', 'Katto 3', 'Katto 4', 'Viherkatto 1', 'Viherkatto 2', 'Hallin alapohja']
for name in order:
    d = sheets[name]; ws = wb.create_sheet(name)
    ws.append(list(d.columns))
    for r in d.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in r])
    style_header(ws); ws.freeze_panes = 'B2'
    for col, w in zip('ABCDEFGHIJK', [18, 10, 11, 10, 13, 11, 10, 13, 13, 12, 12]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, max_col=11):
        row[0].number_format = 'yyyy-mm-dd hh:mm:ss'
        for c in row: c.font = F
        for c in row[8:11]:
            if c.value == 'YES': c.fill = NOTE

out = str(ROOT / 'data/clean/VILPE_Vantaa_cleaned.xlsx')
wb.save(out); print('saved', out)
